#!/usr/bin/env python3
"""
Import a Phase 1 workbook (.xlsx) into the Phase 2 canonical database.

Usage:
  python3 scripts/import_phase1.py --workbook workbook/group-dynamics-simulator-phase1.xlsx --preflight-only
  python3 scripts/import_phase1.py --workbook workbook/group-dynamics-simulator-phase1.xlsx
"""

from __future__ import annotations

import argparse
import os
import sys
from dataclasses import dataclass, field
from typing import Any
import re

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from app.database import SessionLocal, init_db
from app.models.entities import (
    AssessmentSnapshot,
    GroupContext,
    GroupMembership,
    Person,
    RelationshipEdge,
    Scenario,
    SimulationConfig,
)
from app.services.validation import preflight_assessment


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in {"true", "1", "yes", "y"}


def _as_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    return int(value)


def _as_list(value: Any) -> list[str] | None:
    if value is None:
        return None
    lines = [line.strip("- ").strip() for line in str(value).splitlines() if line.strip()]
    return lines or None


def _sheet_rows(ws, min_row: int = 2, min_col: int = 1, max_col: int | None = None):
    for row in ws.iter_rows(min_row=min_row, min_col=min_col, max_col=max_col, values_only=True):
        yield row


_PERSON_ID_RE = re.compile(r"^[a-z0-9._-]{3,64}$")


def _looks_like_person_id(value: Any) -> bool:
    if value is None:
        return False
    return bool(_PERSON_ID_RE.match(str(value).strip()))


@dataclass
class PreflightResult:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def ok(self) -> bool:
        return len(self.errors) == 0


def parse_people(wb) -> list[dict[str, Any]]:
    ws = wb["People"]
    people = []
    for person_id, display_name, role, group_membership, authority_level, is_active, *_ in _sheet_rows(ws):
        if not _looks_like_person_id(person_id):
            continue
        people.append(
            {
                "id": str(person_id).strip(),
                "display_name": str(display_name).strip(),
                "role": str(role).strip(),
                "group_membership": str(group_membership).strip(),
                "authority_level": _as_int(authority_level),
                "is_active": _as_bool(is_active),
            }
        )
    return people


def parse_assessments(wb) -> dict[str, dict[str, Any]]:
    snapshots: dict[str, dict[str, Any]] = {}

    def ensure(person_id: str) -> dict[str, Any]:
        if person_id not in snapshots:
            snapshots[person_id] = {"person_id": person_id}
        return snapshots[person_id]

    for person_id, o, c, e, a, n, evidence_source, *_ in _sheet_rows(wb["Big Five"]):
        if not _looks_like_person_id(person_id):
            continue
        snap = ensure(str(person_id).strip())
        snap.update(
            {
                "big_five_openness": _as_int(o),
                "big_five_conscientiousness": _as_int(c),
                "big_five_extraversion": _as_int(e),
                "big_five_agreeableness": _as_int(a),
                "big_five_neuroticism": _as_int(n),
                "evidence_source": str(evidence_source or "self_report"),
            }
        )

    for person_id, comp, collab, comprom, avoid, accom, *_tail in _sheet_rows(wb["Conflict Style"]):
        if not _looks_like_person_id(person_id):
            continue
        snap = ensure(str(person_id).strip())
        snap.update(
            {
                "conflict_competing": _as_int(comp),
                "conflict_collaborating": _as_int(collab),
                "conflict_compromising": _as_int(comprom),
                "conflict_avoiding": _as_int(avoid),
                "conflict_accommodating": _as_int(accom),
            }
        )

    for person_id, i1, i2, i3, i4, i5, i6, i7, *_tail in _sheet_rows(wb["Psych Safety"]):
        if not _looks_like_person_id(person_id):
            continue
        snap = ensure(str(person_id).strip())
        snap.update(
            {
                "psych_safety_item_1": _as_int(i1),
                "psych_safety_item_2": _as_int(i2),
                "psych_safety_item_3": _as_int(i3),
                "psych_safety_item_4": _as_int(i4),
                "psych_safety_item_5": _as_int(i5),
                "psych_safety_item_6": _as_int(i6),
                "psych_safety_item_7": _as_int(i7),
            }
        )

    for person_id, d, co, vd, lq, ft, avs, ra, ds, at, *_tail in _sheet_rows(wb["Comm-Decision"]):
        if not _looks_like_person_id(person_id):
            continue
        snap = ensure(str(person_id).strip())
        snap.update(
            {
                "comm_directness": _as_int(d),
                "comm_context_orientation": _as_int(co),
                "comm_verbal_dominance": _as_int(vd),
                "comm_listening_quality": _as_int(lq),
                "comm_feedback_tolerance": _as_int(ft),
                "decision_analytical_vs_intuitive": _as_int(avs),
                "decision_risk_appetite": _as_int(ra),
                "decision_speed": _as_int(ds),
                "decision_ambiguity_tolerance": _as_int(at),
            }
        )

    for person_id, p, u, und, m, *_tail in _sheet_rows(wb["EQ"]):
        if not _looks_like_person_id(person_id):
            continue
        snap = ensure(str(person_id).strip())
        snap.update(
            {
                "eq_perceiving": _as_int(p),
                "eq_using": _as_int(u),
                "eq_understanding": _as_int(und),
                "eq_managing": _as_int(m),
            }
        )

    for person_id, sec, anx, avo, fear, *_tail in _sheet_rows(wb["Attachment"]):
        if not _looks_like_person_id(person_id):
            continue
        snap = ensure(str(person_id).strip())
        snap.update(
            {
                "attachment_secure": _as_int(sec),
                "attachment_anxious": _as_int(anx),
                "attachment_avoidant": _as_int(avo),
                "attachment_fearful": _as_int(fear),
            }
        )

    return snapshots


def parse_relationships(wb) -> list[dict[str, Any]]:
    ws = wb["Relationship Matrix"]
    edges = []
    for row in _sheet_rows(ws):
        from_id, to_id = row[0], row[1]
        if not from_id or not to_id:
            continue
        edges.append(
            {
                "from_person_id": str(from_id).strip(),
                "to_person_id": str(to_id).strip(),
                "trust": _as_int(row[2]),
                "influence": _as_int(row[3]),
                "emotional_closeness": _as_int(row[4]),
                "respect": _as_int(row[5]),
                "conflict_intensity": _as_int(row[6]),
                "dependency": _as_int(row[7]),
                "communication_frequency": _as_int(row[8]),
                "avoidance": _as_int(row[9]),
                "alliance": _as_int(row[10]),
                "power_differential": _as_int(row[11]),
                "evidence_source": str(row[13] or "self_report"),
                "notes": str(row[14]).strip() if row[14] else None,
            }
        )
    return edges


def parse_key_value_sheet(wb, sheet_name: str) -> dict[str, Any]:
    ws = wb[sheet_name]
    values = {}
    for key, value, *_ in _sheet_rows(ws):
        if not key:
            continue
        values[str(key).strip()] = value
    return values


def preflight(people: list[dict], assessments: dict[str, dict], edges: list[dict], group: dict, scenario: dict, config: dict) -> PreflightResult:
    result = PreflightResult()

    people_ids = {p["id"] for p in people}
    required_person = ["id", "display_name", "role", "group_membership", "authority_level", "is_active"]
    for person in people:
        for field_name in required_person:
            if person.get(field_name) in (None, ""):
                result.errors.append(f"people[{person.get('id','<missing>')}].{field_name} is required")

    for person_id, snapshot in assessments.items():
        if person_id not in people_ids:
            result.errors.append(f"assessment person_id '{person_id}' not found in People tab")
        for err in preflight_assessment(snapshot):
            result.errors.append(f"assessment[{person_id}] {err}")

    for edge in edges:
        if edge["from_person_id"] == edge["to_person_id"]:
            result.errors.append(f"relationship self-edge detected for '{edge['from_person_id']}'")
        if edge["from_person_id"] not in people_ids or edge["to_person_id"] not in people_ids:
            result.errors.append(
                f"relationship edge ({edge['from_person_id']} -> {edge['to_person_id']}) references unknown person"
            )

    required_group = ["group.id", "group.type", "group.structure", "group.shared_goals", "group.stress_level"]
    for key in required_group:
        if group.get(key) in (None, ""):
            result.errors.append(f"group field '{key}' is required")

    required_scenario = [
        "scenario.id", "scenario.title", "scenario.type", "scenario.trigger_event",
        "scenario.stakes_level", "scenario.emotional_intensity", "scenario.ambiguity_level",
        "scenario.time_pressure", "scenario.required_decision", "scenario.success_criteria",
        "scenario.failure_consequences",
    ]
    for key in required_scenario:
        if scenario.get(key) in (None, ""):
            result.errors.append(f"scenario field '{key}' is required")

    required_config = [
        "sim.prompt_version_key", "sim.passes", "sim.randomness", "sim.depth",
        "sim.dialogue_enabled", "sim.report_detail_level", "sim.intervention_mode",
        "sim.evidence_strictness", "sim.guardrail_verbosity",
    ]
    for key in required_config:
        if config.get(key) in (None, ""):
            result.errors.append(f"config field '{key}' is required")

    if not assessments:
        result.warnings.append("No assessment rows found.")
    if not edges:
        result.warnings.append("No relationship edges found.")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Import Phase 1 workbook into Phase 2 DB")
    parser.add_argument("--workbook", required=True, help="Path to .xlsx workbook")
    parser.add_argument("--preflight-only", action="store_true", help="Validate workbook and exit without writing DB")
    args = parser.parse_args()

    wb = load_workbook(args.workbook, data_only=True)

    people = parse_people(wb)
    assessments = parse_assessments(wb)
    edges = parse_relationships(wb)
    group = parse_key_value_sheet(wb, "Group Context")
    scenario = parse_key_value_sheet(wb, "Scenario Builder")
    config = parse_key_value_sheet(wb, "Simulation Config")

    report = preflight(people, assessments, edges, group, scenario, config)
    if report.warnings:
        print("Preflight warnings:")
        for warning in report.warnings:
            print(f"  - {warning}")

    if not report.ok():
        print("Preflight errors:")
        for err in report.errors:
            print(f"  - {err}")
        return 1

    print(
        f"Preflight passed: people={len(people)}, assessments={len(assessments)}, "
        f"relationships={len(edges)}"
    )
    if args.preflight_only:
        return 0

    init_db()
    db = SessionLocal()
    try:
        # People
        for payload in people:
            existing = db.get(Person, payload["id"])
            if existing:
                for k, v in payload.items():
                    setattr(existing, k, v)
            else:
                db.add(Person(**payload))

        # Group
        group_id = str(group["group.id"])
        group_payload = {
            "id": group_id,
            "type": str(group["group.type"]),
            "structure": str(group["group.structure"]),
            "shared_goals": str(group["group.shared_goals"]),
            "norms_explicit": group.get("group.norms.explicit"),
            "norms_implicit": group.get("group.norms.implicit"),
            "decision_rules": group.get("group.decision_rules"),
            "conflict_history": group.get("group.conflict_history"),
            "stress_level": _as_int(group.get("group.stress_level")),
            "role_clarity": _as_int(group.get("group.role_clarity")),
            "cultural_context": group.get("group.cultural_context"),
            "environmental_constraints": group.get("group.environmental_constraints"),
        }
        existing_group = db.get(GroupContext, group_id)
        if existing_group:
            for k, v in group_payload.items():
                setattr(existing_group, k, v)
        else:
            db.add(GroupContext(**group_payload))

        # Memberships (idempotent)
        for p in people:
            membership = (
                db.query(GroupMembership)
                .filter(GroupMembership.group_id == group_id, GroupMembership.person_id == p["id"])
                .first()
            )
            if not membership:
                db.add(GroupMembership(group_id=group_id, person_id=p["id"]))

        # Scenario
        scenario_id = str(scenario["scenario.id"])
        scenario_payload = {
            "id": scenario_id,
            "title": str(scenario["scenario.title"]),
            "type": str(scenario["scenario.type"]),
            "trigger_event": str(scenario["scenario.trigger_event"]),
            "stakes_level": _as_int(scenario.get("scenario.stakes_level")),
            "emotional_intensity": _as_int(scenario.get("scenario.emotional_intensity")),
            "ambiguity_level": _as_int(scenario.get("scenario.ambiguity_level")),
            "time_pressure": _as_int(scenario.get("scenario.time_pressure")),
            "resource_constraints": scenario.get("scenario.resource_constraints"),
            "public_visibility": _as_bool(scenario.get("scenario.public_visibility")),
            "required_decision": str(scenario["scenario.required_decision"]),
            "success_criteria": str(scenario["scenario.success_criteria"]),
            "failure_consequences": str(scenario["scenario.failure_consequences"]),
            "known_facts": _as_list(scenario.get("scenario.known_facts")),
            "uncertain_facts": _as_list(scenario.get("scenario.uncertain_facts")),
            "intervention_options": _as_list(scenario.get("scenario.intervention_options")),
            "group_id": group_id,
        }
        existing_scenario = db.get(Scenario, scenario_id)
        if existing_scenario:
            for k, v in scenario_payload.items():
                setattr(existing_scenario, k, v)
        else:
            db.add(Scenario(**scenario_payload))

        # Config
        config_payload = {
            "prompt_version_key": str(config["sim.prompt_version_key"]),
            "passes": _as_int(config["sim.passes"]),
            "randomness": str(config["sim.randomness"]),
            "depth": str(config["sim.depth"]),
            "dialogue_enabled": _as_bool(config["sim.dialogue_enabled"]),
            "report_detail_level": str(config["sim.report_detail_level"]),
            "intervention_mode": str(config["sim.intervention_mode"]),
            "evidence_strictness": str(config["sim.evidence_strictness"]),
            "guardrail_verbosity": str(config["sim.guardrail_verbosity"]),
        }
        db.add(SimulationConfig(**config_payload))

        # Assessments (replace existing snapshot for this import)
        for person_id, payload in assessments.items():
            existing = db.query(AssessmentSnapshot).filter(AssessmentSnapshot.person_id == person_id).first()
            if existing:
                for k, v in payload.items():
                    if k != "person_id":
                        setattr(existing, k, v)
            else:
                db.add(AssessmentSnapshot(**payload))

        # Relationship edges
        for payload in edges:
            existing = (
                db.query(RelationshipEdge)
                .filter(
                    RelationshipEdge.from_person_id == payload["from_person_id"],
                    RelationshipEdge.to_person_id == payload["to_person_id"],
                )
                .first()
            )
            if existing:
                for k, v in payload.items():
                    setattr(existing, k, v)
            else:
                db.add(RelationshipEdge(**payload))

        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    print("Import complete.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
